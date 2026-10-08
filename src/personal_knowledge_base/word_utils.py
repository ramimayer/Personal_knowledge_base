from docx import Document
from pathlib import Path
import os

import win32com.client as win32
from win32com.client import constants
import pythoncom

import re

#----------------------------------------------------------------------------------------
def docx_to_txt(file: str, filename: str, output_folder: str):
    Path(output_folder).mkdir(parents=True, exist_ok=True)
    output_file_path = output_folder + "/" + filename + ".txt"
    
    # Check if the output file already exists
    if os.path.exists(output_file_path):
        return

    text = load_docx_table(file)
    text += load_docx_text(file)

    with open(output_file_path, "w") as f:
        f.write(text)

#----------------------------------------------------------------------------------------
def doc_to_txt(file: str, filename: str, output_folder: str):
    print("doc_to_txt(): input_file_path = " + file + ", output_folder = " + output_folder)
    # Path(output_folder).mkdir(parents=True, exist_ok=True)
    # output_file_path = output_folder + "/" + filename + ".txt"
    
    # # Check if the output file already exists
    # if os.path.exists(output_file_path):
    #     print("doc_to_txt(): output_file_path already exists = " + output_file_path)
    #     return

    # """
    # # substitute all the slashes in the file path with double backslashes"""

    # file = re.sub(r"/", "\\", file)
    # file_windows = file.replace("\\", "\\\\")
    # print("file_windows = " + file)

    # my_file_windows = "C:\\Users\\ramim\\Documents\\projects\\data\\docs_orig_format\\Memory_Items.doc"
    # print("file_windows = " + file_windows)
    # docx_file = save_as_docx(file_windows)
    # docx_to_txt(docx_file, filename, output_folder)
    # Path(output_folder).unlink(docx_file)  # Delete the temporary .docx file after conversion

#----------------------------------------------------------------------------------------
def save_as_docx(doc_file:str) -> str:
    pythoncom.CoInitialize()  # Initialize the COM library for the current thread
    # Opening MS Word
    print("save_as_docx(): doc_file = " + doc_file)
    word = win32.gencache.EnsureDispatch('Word.Application')
    doc = word.Documents.Open(doc_file)
    print("bfr doc.Activate()")
    doc.Activate ()
    print("aftr doc.Activate()")

    new_file_abs = doc_file
    new_file_abs = re.sub(r'\.\w+$', '.docx', new_file_abs)
    print(new_file_abs)

    word.ActiveDocument.SaveAs(
        new_file_abs, FileFormat=constants.wdFormatXMLDocument
    )
    doc.Close(False)

    return new_file_abs

#----------------------------------------------------------------------------------------
def load_docx_text(file: str) -> str:
    text = ""
    document = Document(file)
    for p in document.paragraphs:
        text += p.text + "\n"

    return text

#----------------------------------------------------------------------------------------
def load_docx_table(file: str) -> str:
    text = ""
    # Data will be a list of rows represented as dictionaries
    # containing each row's data.
    l = []
    document = Document(file)
    if (len(document.tables) > 0):
        print(len(document.tables))
        table = document.tables[0]
        
        keys = None
        for i, row in enumerate(table.rows):
            t = (cell.text for cell in row.cells)

            # Establish the mapping based on the first row
            # headers; these will become the keys of our dictionary
            if i == 0:
                keys = tuple(t)
                continue

            # Construct a dictionary for this row, mapping
            # keys to values for this row
            row_data = dict(zip(keys, t))
            l.append(row_data)

        for d in l:
            for k, v in d.items():
                text += k + ": " + v + ".      "
            text += "\n\n"

    return text